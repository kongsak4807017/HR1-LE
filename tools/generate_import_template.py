from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

AGE_GROUPS = ["0", "1-4"] + [f"{a}-{a+4}" for a in range(5, 85, 5)] + ["85+"]

def style_header(ws, row=1):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def build(path="hr1_le_import_template.xlsx"):
    wb = Workbook()

    readme = wb.active
    readme.title = "README"
    readme.append(["HR1-LE Canonical Import Template", ""])
    readme.append(["Scope", "Health Region 1: 8 provinces"])
    readme.append(["Important", "Compatibility schema; exact target proprietary headers remain to be verified."])
    readme.column_dimensions["A"].width = 22
    readme.column_dimensions["B"].width = 90

    metadata = wb.create_sheet("metadata")
    metadata.append([
        "dataset_name","dataset_year_start","dataset_year_end","area_scope",
        "population_source","death_source","year_calendar","created_by",
        "method_version","notes"
    ])
    metadata.append(["HR1 mortality analytics",2020,2025,"region","","","CE","","",""])
    style_header(metadata)

    population = wb.create_sheet("population")
    population.append(["year","area_code","area_name","area_level","sex","age_group","population","source"])
    population.append([2025,"50","เชียงใหม่","province","M","0",0,""])
    style_header(population)

    deaths = wb.create_sheet("deaths")
    deaths.append(["year","area_code","area_name","area_level","sex","age_group","icd10","deaths","source"])
    deaths.append([2025,"50","เชียงใหม่","province","M","65-69","I21",0,""])
    style_header(deaths)

    cause = wb.create_sheet("cause_mapping")
    cause.append(["icd10","cause_group_code","cause_name_th","cause_name_en","mapping_version"])
    style_header(cause)

    standard = wb.create_sheet("standard_le")
    standard.append(["reference_name","age_group","remaining_le"])
    for age in AGE_GROUPS:
        standard.append(["approved-reference", age, None])
    style_header(standard)

    priority = wb.create_sheet("priority_scores")
    priority.append([
        "round_id","rater_id","disease_code","disease_name","problem_size_basis",
        "rate_per_100k","problem_size_score","severity_score","epidemic_score",
        "social_economic_score","feasibility_score","health_gain_score",
        "public_perception_score","total_score"
    ])
    style_header(priority)

    lists = wb.create_sheet("lists")
    lists.append(["age_group","sex","area_level","problem_size_basis"])
    max_len = max(len(AGE_GROUPS), 4)
    for i in range(max_len):
        lists.append([
            AGE_GROUPS[i] if i < len(AGE_GROUPS) else None,
            ["M","F"][i] if i < 2 else None,
            ["region","province","district","subdistrict"][i] if i < 4 else None,
            ["incidence","prevalence","mortality","case_fatality"][i] if i < 4 else None,
        ])
    style_header(lists)

    dv_sex = DataValidation(type="list", formula1='"M,F"', allow_blank=False)
    dv_level = DataValidation(type="list", formula1='"region,province,district,subdistrict"', allow_blank=False)
    for ws in (population, deaths):
        ws.add_data_validation(dv_sex)
        ws.add_data_validation(dv_level)
        dv_level.add(f"D2:D10000")
        dv_sex.add(f"E2:E10000")
        ws.freeze_panes = "A2"

    for ws in wb.worksheets:
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = min(max(12, max(len(str(c.value or "")) for c in col) + 2), 40)

    wb.save(path)
    return path

if __name__ == "__main__":
    print(build())
